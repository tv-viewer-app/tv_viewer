from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "templates" / "lg-self-evaluation-checklist-5.0-original.xlsx"
OUTPUT = ROOT / "lg-self-evaluation-checklist-5.0-tv-viewer-draft.xlsx"

PASS = {
    5: "Headless Chromium launch and packaged-app smoke test passed.",
    6: "Main page, loading state, splash, overscan-safe layout, and 16:9 screenshots reviewed at 1920x1080.",
    9: "Both 1920x1080 and 1280x720 IPKs are built and validated.",
    10: "English UI and Seller Lounge screenshots visually reviewed for clipping and broken text.",
    11: "All selectable controls implement a visible white focus outline and scale state.",
    14: "Automated browser flow exercised notice, filters, Help, channel selection, and player controls.",
    15: "On-screen Back to channels control returns from player to catalog.",
    26: "Search and category filtering exercised against the live curated catalog.",
}

NOT_APPLICABLE = {
    8: "No advertising is included.",
    16: "No dedicated Exit UI button is included; platform Back behavior is used.",
    21: "Voice search is not implemented.",
    22: "No account terms or membership flow; only an informational first-run notice.",
    23: "No sign-up capability.",
    24: "No sign-in capability.",
    25: "No sign-out capability.",
    27: "Adult content categories are excluded and explicit names are blocked.",
    28: "No payment or in-app purchase capability.",
    29: "No purchased-content flow.",
    36: "Color-key shortcuts are not implemented.",
    38: "The app is not Magic-Remote-only.",
    39: "The app is not Magic-Remote-only.",
    43: "LIVE key behavior is not app-defined.",
    45: "Version 2.25.0 supports English only and has no in-app language switcher.",
    48: "The player uses a full-screen presentation and has no full/original toggle.",
    50: "Catalog content is live streaming rather than episodic replay content.",
    52: "Subtitle controls are not implemented.",
    53: "Resume playback is not applicable to live streams.",
    56: "No paid or DRM-protected content is offered.",
}

PENDING = (
    "PENDING: Must be executed on a physical LG TV or webOS Cloud Test Lab "
    "before changing this cell to Pass."
)

wb = load_workbook(SOURCE)
ws = wb["Self-Checklist_Input Required"]
for row in range(5, 58):
    result = ws.cell(row, 14)
    comment = ws.cell(row, 15)
    if row in PASS:
        result.value = "Pass"
        comment.value = PASS[row]
    elif row in NOT_APPLICABLE:
        result.value = "N/A"
        comment.value = NOT_APPLICABLE[row]
    else:
        result.value = None
        comment.value = PENDING
        result.fill = PatternFill("solid", fgColor="FFF2CC")
        comment.fill = PatternFill("solid", fgColor="FFF2CC")

other = wb["Other Check Points"]
static_pass = {
    4: "TV Viewer accurately describes the product and is not a misleading generic category name.",
    5: "Title contains no inappropriate wording.",
    7: "Seller is identified as TV Viewer, not LG.",
    9: "400x400 store icon and packaged icons are square.",
    10: "Icon artwork uses rounded inner artwork within the required square canvas.",
    11: "Store icon has an opaque background.",
    12: "Generated screenshots contain no sexual or violent imagery.",
    13: "All submitted screenshots are distinct.",
    14: "Description matches catalog, search, favorites, remote navigation, and playback functionality.",
    15: "Description exceeds the minimum and explains the application clearly.",
    16: "Description does not imply that LG is the seller.",
    17: "1920x1080 non-black splash image and required metadata are packaged.",
    18: "English title, description, and UI are provided. Territory selection remains pending.",
}
for row in range(4, 55):
    result = other.cell(row, 7)
    comment = other.cell(row, 8)
    if row in static_pass:
        result.value = "Pass"
        comment.value = static_pass[row]
    elif row in range(46, 53):
        result.value = "N/A"
        comment.value = "The app contains no gaming, game money, gambling, or rewards system."
    else:
        result.value = None
        comment.value = (
            "PENDING: Requires final manual content sample and rights review of "
            "the curated live catalog before submission."
        )
        result.fill = PatternFill("solid", fgColor="FFF2CC")
        comment.fill = PatternFill("solid", fgColor="FFF2CC")

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = copy(sheet.freeze_panes)

ws["N1"].comment = Comment(
    "DRAFT prepared September 30, 2026. Yellow blank result cells require "
    "physical LG TV or webOS Cloud Test Lab verification before upload.",
    "TV Viewer",
)
wb.save(OUTPUT)
print(OUTPUT)
